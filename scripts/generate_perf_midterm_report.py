#!/usr/bin/env python3
from __future__ import annotations

import json
import math
import statistics
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
RESULTS_ROOT = ROOT / 'perf' / 'results'
OUTPUT_ROOT = ROOT / 'output' / 'doc'
ASSET_ROOT = OUTPUT_ROOT / 'assets'


@dataclass
class RunMetrics:
    scenario: str
    file_name: str
    avg_ms: float
    p95_ms: float
    p99_ms: float
    rps: float
    error_rate: float
    req_count: int
    incomplete: bool


@dataclass
class AggregateMetrics:
    scenario: str
    n: int
    avg_ms: float
    avg_ms_ci95: float
    p95_ms: float
    p95_ms_ci95: float
    p99_ms: float
    p99_ms_ci95: float
    rps: float
    rps_ci95: float
    error_rate: float
    error_rate_ci95: float


@dataclass
class HostSpecs:
    model_name: str
    model_identifier: str
    chip: str
    cpu_cores: str
    memory_gb: str
    macos_version: str
    kernel_version: str


def safe_metric(metrics: dict, metric_name: str, stat: str, default=0):
    value = metrics.get(metric_name, {}).get(stat, default)
    return default if value is None else value


def is_incomplete_run(metrics: dict) -> bool:
    req_count = safe_metric(metrics, 'http_reqs', 'count', 0)
    sent_count = safe_metric(metrics, 'data_sent', 'count', 0)
    received_count = safe_metric(metrics, 'data_received', 'count', 0)
    duration_avg = safe_metric(metrics, 'http_req_duration', 'avg', 0)
    duration_max = safe_metric(metrics, 'http_req_duration', 'max', 0)
    return (
        req_count > 0
        and sent_count == 0
        and received_count == 0
        and duration_avg == 0
        and duration_max == 0
    )


def load_run(path: Path, scenario: str) -> RunMetrics:
    data = json.loads(path.read_text())
    metrics = data['metrics']
    return RunMetrics(
        scenario=scenario,
        file_name=path.name,
        avg_ms=float(safe_metric(metrics, 'http_req_duration', 'avg', 0)),
        p95_ms=float(safe_metric(metrics, 'http_req_duration', 'p(95)', 0)),
        p99_ms=float(safe_metric(metrics, 'http_req_duration', 'p(99)', 0)),
        rps=float(safe_metric(metrics, 'http_reqs', 'rate', 0)),
        error_rate=float(safe_metric(metrics, 'http_req_failed', 'value', 0)),
        req_count=int(safe_metric(metrics, 'http_reqs', 'count', 0)),
        incomplete=is_incomplete_run(metrics),
    )


def ci95(values: list[float]) -> float:
    if len(values) <= 1:
        return 0.0
    return 1.96 * statistics.stdev(values) / math.sqrt(len(values))


def aggregate_runs(scenario: str, runs: Iterable[RunMetrics]) -> AggregateMetrics:
    valid_runs = [run for run in runs if not run.incomplete]
    if not valid_runs:
        raise ValueError(f'No valid runs for scenario {scenario}')

    def vals(attr: str) -> list[float]:
        return [float(getattr(run, attr)) for run in valid_runs]

    avg_vals = vals('avg_ms')
    p95_vals = vals('p95_ms')
    p99_vals = vals('p99_ms')
    rps_vals = vals('rps')
    err_vals = vals('error_rate')

    return AggregateMetrics(
        scenario=scenario,
        n=len(valid_runs),
        avg_ms=statistics.mean(avg_vals),
        avg_ms_ci95=ci95(avg_vals),
        p95_ms=statistics.mean(p95_vals),
        p95_ms_ci95=ci95(p95_vals),
        p99_ms=statistics.mean(p99_vals),
        p99_ms_ci95=ci95(p99_vals),
        rps=statistics.mean(rps_vals),
        rps_ci95=ci95(rps_vals),
        error_rate=statistics.mean(err_vals),
        error_rate_ci95=ci95(err_vals),
    )


def fmt(value: float, digits: int = 2) -> str:
    return f'{value:.{digits}f}'


def ensure_output_dirs() -> None:
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    ASSET_ROOT.mkdir(parents=True, exist_ok=True)


def read_command_output(command: list[str]) -> str:
    try:
        completed = subprocess.run(
            command,
            check=True,
            capture_output=True,
            text=True,
        )
    except Exception:
        return ''
    return completed.stdout


def get_host_specs() -> HostSpecs:
    hardware_text = read_command_output(['system_profiler', 'SPHardwareDataType'])
    software_text = read_command_output(['system_profiler', 'SPSoftwareDataType'])

    def extract(label: str, text: str, fallback: str = 'unknown') -> str:
        prefix = f'{label}:'
        for line in text.splitlines():
            stripped = line.strip()
            if stripped.startswith(prefix):
                return stripped.split(':', 1)[1].strip()
        return fallback

    return HostSpecs(
        model_name=extract('Model Name', hardware_text),
        model_identifier=extract('Model Identifier', hardware_text),
        chip=extract('Chip', hardware_text),
        cpu_cores=extract('Total Number of Cores', hardware_text),
        memory_gb=extract('Memory', hardware_text),
        macos_version=extract('System Version', software_text),
        kernel_version=extract('Kernel Version', software_text),
    )


def load_dataset() -> tuple[list[RunMetrics], list[RunMetrics], dict[int, list[RunMetrics]], dict[int, AggregateMetrics]]:
    baseline_runs = [
        load_run(path, 'baseline')
        for path in sorted((RESULTS_ROOT / 'baseline').glob('*.json'))
    ]
    baseline_valid = [run for run in baseline_runs if not run.incomplete]

    load_runs: dict[int, list[RunMetrics]] = {}
    aggregates: dict[int, AggregateMetrics] = {}
    for vus in [10, 50, 100, 200, 500]:
        scenario = f'load-scaling/light/vus-{vus}'
        run_paths = sorted((RESULTS_ROOT / 'load-scaling' / 'light' / f'vus-{vus}').glob('run-*.json'))
        runs = [load_run(path, scenario) for path in run_paths]
        load_runs[vus] = runs
        aggregates[vus] = aggregate_runs(scenario, runs)

    return baseline_runs, baseline_valid, load_runs, aggregates


def build_latency_chart(aggregates: dict[int, AggregateMetrics]) -> Path:
    x = list(aggregates.keys())
    avg = [aggregates[v].avg_ms for v in x]
    p95 = [aggregates[v].p95_ms for v in x]
    p99 = [aggregates[v].p99_ms for v in x]

    plt.figure(figsize=(8, 4.8))
    plt.plot(x, avg, marker='o', linewidth=2, label='Mean latency')
    plt.plot(x, p95, marker='o', linewidth=2, label='P95 latency')
    plt.plot(x, p99, marker='o', linewidth=2, label='P99 latency')
    plt.yscale('log')
    plt.xticks(x, [str(v) for v in x])
    plt.xlabel('Concurrent virtual users (VUs)')
    plt.ylabel('Latency (ms, log scale)')
    plt.title('Light workload latency vs concurrency')
    plt.grid(True, which='both', linestyle='--', linewidth=0.5, alpha=0.5)
    plt.legend()
    plt.tight_layout()

    path = ASSET_ROOT / 'light_latency_vs_vus.png'
    plt.savefig(path, dpi=200)
    plt.close()
    return path


def build_throughput_error_chart(aggregates: dict[int, AggregateMetrics]) -> Path:
    x = list(aggregates.keys())
    rps = [aggregates[v].rps for v in x]
    err = [aggregates[v].error_rate * 100.0 for v in x]

    fig, ax1 = plt.subplots(figsize=(8, 4.8))
    ax1.plot(x, rps, color='#0b5ed7', marker='o', linewidth=2)
    ax1.set_xlabel('Concurrent virtual users (VUs)')
    ax1.set_ylabel('Throughput (req/s)', color='#0b5ed7')
    ax1.tick_params(axis='y', labelcolor='#0b5ed7')
    ax1.set_xticks(x, [str(v) for v in x])
    ax1.grid(True, linestyle='--', linewidth=0.5, alpha=0.5)

    ax2 = ax1.twinx()
    ax2.plot(x, err, color='#c1121f', marker='s', linewidth=2)
    ax2.set_ylabel('Error rate (%)', color='#c1121f')
    ax2.tick_params(axis='y', labelcolor='#c1121f')

    plt.title('Light workload throughput and error rate vs concurrency')
    plt.tight_layout()

    path = ASSET_ROOT / 'light_throughput_error_vs_vus.png'
    plt.savefig(path, dpi=200)
    plt.close()
    return path


def build_high_load_run_chart(load_runs: dict[int, list[RunMetrics]]) -> Path:
    labels = []
    p95 = []
    error_pct = []
    colors = []

    for vus in [200, 500]:
        for run in load_runs[vus]:
            labels.append(f'{vus}-{run.file_name.replace("run-", "r").replace(".json", "")}')
            p95.append(run.p95_ms)
            error_pct.append(run.error_rate * 100.0)
            colors.append('#6c757d' if run.incomplete else ('#c1121f' if run.error_rate >= 0.05 else '#0b5ed7'))

    fig, ax1 = plt.subplots(figsize=(10, 5))
    ax1.bar(labels, p95, color=colors, alpha=0.8)
    ax1.set_ylabel('P95 latency (ms)')
    ax1.set_xlabel('Run')
    ax1.tick_params(axis='x', rotation=45)
    ax1.grid(True, axis='y', linestyle='--', linewidth=0.5, alpha=0.5)

    ax2 = ax1.twinx()
    ax2.plot(labels, error_pct, color='#111111', marker='o', linewidth=1.8)
    ax2.set_ylabel('Error rate (%)')

    plt.title('High-load run variability for the light workload')
    plt.tight_layout()

    path = ASSET_ROOT / 'light_high_load_run_variability.png'
    plt.savefig(path, dpi=200)
    plt.close()
    return path


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ''
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def add_caption(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True)
        shade_cell(table.rows[0].cells[idx], 'D9EAF7')
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    document.add_paragraph()


def write_markdown(
    report_path: Path,
    baseline_valid: list[RunMetrics],
    load_runs: dict[int, list[RunMetrics]],
    aggregates: dict[int, AggregateMetrics],
    host_specs: HostSpecs,
) -> None:
    baseline_avg = statistics.mean(run.avg_ms for run in baseline_valid)
    baseline_p95 = statistics.mean(run.p95_ms for run in baseline_valid)
    baseline_rps = statistics.mean(run.rps for run in baseline_valid)

    valid_500 = [run for run in load_runs[500] if not run.incomplete]
    incomplete_500 = [run for run in load_runs[500] if run.incomplete]

    lines = [
        '# P.I.T.C.H. Midterm Performance Evaluation Report',
        '',
        'Date: 2026-03-15',
        '',
        '## Abstract',
        '',
        'This midterm report presents the current performance evaluation progress for P.I.T.C.H., a microservice-based AI-assisted web application. The current dataset covers baseline characterization and light API load-scaling experiments at 10, 50, 100, 200, and 500 concurrent virtual users using k6. Initial results show that the light endpoint performs well at low load, but tail latency degrades sharply by 50 and 100 virtual users. At 200 virtual users the system becomes unstable, and at 500 virtual users the gateway enters overload behavior with partial and fully incomplete runs. These results provide a defensible midterm answer to the first project objective and identify the gateway as the leading bottleneck candidate for the final report, pending confirmation from service-level observability data.',
        '',
        '## I. Introduction',
        '',
        'P.I.T.C.H. is a cloud-native microservice system composed of a gateway, multiple NestJS backend services, asynchronous messaging over RabbitMQ, Redis caching, persistent databases, and a Next.js frontend. The goal of this project is to evaluate how the system behaves as concurrent demand increases, determine where bottlenecks emerge, and quantify the impact of future scaling decisions.',
        '',
        '## II. Background and Motivation',
        '',
        'Microservice architectures offer deployment flexibility and independent scaling, but they also introduce network, queuing, coordination, and observability complexity. For P.I.T.C.H., performance matters directly to user experience because requests traverse a gateway, backend services, and messaging infrastructure. The motivation for this study is to move from architecture assumptions to measured evidence and use that evidence to identify the system limits and the most valuable optimization targets.',
        '',
        '## III. Performance Evaluation Objectives',
        '',
        '1. Measure how response time and throughput change as concurrent load increases.',
        '2. Identify which service or dependency becomes the bottleneck under stress.',
        '3. Quantify how scaling and caching affect latency, throughput, and stability in later experiments.',
        '',
        '## IV. Evaluation Technique',
        '',
        'The current evaluation uses quantitative, measurement-based testing. Closed-workload load generation is performed with k6, and summary metrics are exported to JSON after each run. The measurements emphasized in the current dataset are mean latency, P95 latency, P99 latency, throughput, and error rate. Warm-up, steady-state, and cool-down phases are included in each run to reduce startup bias.',
        '',
        '## V. Evaluation Setup',
        '',
        'The system under test is deployed in Docker containers. Current collected runs target the gateway entry point and exercise the light API path through the microservice stack. The broader observability environment includes Prometheus, Grafana, RabbitMQ, Redis, container monitoring, and application metrics, though this report focuses on the k6 results collected so far.',
        '',
        'Host machine used for data collection:',
        '',
        f'- model: {host_specs.model_name} ({host_specs.model_identifier})',
        f'- processor: {host_specs.chip}',
        f'- CPU cores: {host_specs.cpu_cores}',
        f'- memory: {host_specs.memory_gb}',
        f'- operating system: {host_specs.macos_version}',
        f'- kernel: {host_specs.kernel_version}',
        '',
        '## VI. Test Workloads',
        '',
        'The current midterm dataset includes:',
        '',
        '- baseline characterization of the light endpoint',
        '- light API load scaling at 10, 50, 100, 200, and 500 virtual users',
        '',
        'AI-intensive, mixed, spike, stress, replica-comparison, and cache-comparison workloads are not yet included in this report and remain part of the final-evaluation plan.',
        '',
        '## VII. Experimental Design',
        '',
        'For the light workload, the target concurrency levels were 10, 50, 100, 200, and 500 virtual users. Five repeated runs were planned per load level. At the highest load levels, threshold failures and gateway instability emerged, so some runs represent overload behavior rather than clean steady-state service. In particular, some of the 500-virtual-user runs became incomplete and must be treated separately during analysis.',
        '',
        '## VIII. Data Collection',
        '',
        'Data was collected from the k6 summary JSON files stored under `perf/results`. Two valid baseline runs were available after excluding one incomplete failed baseline. For load scaling, the light workload dataset now includes 5 runs at 10 VUs, 5 runs at 50 VUs, 4 runs at 100 VUs, 5 runs at 200 VUs, and 5 runs at 500 VUs. Of the 500-VU runs, only 2 are valid complete runs; the remaining 3 are incomplete overload runs with zero-byte transfers and should not be averaged with the valid runs.',
        '',
        'An important threat to validity is that the data was collected on a personal laptop that was also being used for normal student work during the testing window. Background applications, browser activity, course-related development tools, and other foreground tasks were not fully eliminated. As a result, the measured values should be interpreted as realistic development-machine results rather than fully isolated laboratory measurements. The concurrency trends and failure boundaries are still useful, but the absolute latency and throughput numbers may include noise from unrelated host activity.',
        '',
        '## IX. Initial Results',
        '',
        f'Baseline performance was strong, with mean latency {fmt(baseline_avg, 2)} ms, P95 latency {fmt(baseline_p95, 2)} ms, and throughput {fmt(baseline_rps, 2)} req/s. At 10 VUs the system remained effectively identical to baseline. By 50 VUs, throughput increased to {fmt(aggregates[50].rps, 2)} req/s, but P99 latency increased to {fmt(aggregates[50].p99_ms, 2)} ms. At 100 VUs, throughput increased further to {fmt(aggregates[100].rps, 2)} req/s while P95 reached {fmt(aggregates[100].p95_ms, 2)} ms and P99 reached {fmt(aggregates[100].p99_ms, 2)} ms. These results indicate that tail latency degradation begins well before visible request failures.',
        '',
        f'At 200 VUs, the light endpoint entered an instability regime. Across 5 runs, the aggregate error rate reached {fmt(aggregates[200].error_rate * 100, 2)}%, mean latency reached {fmt(aggregates[200].avg_ms, 2)} ms, and P99 latency reached {fmt(aggregates[200].p99_ms, 2)} ms. Run-to-run variability became extreme: one run had nearly {fmt(load_runs[200][3].avg_ms / 1000, 2)} seconds of mean latency and a P99 above {fmt(load_runs[200][3].p99_ms / 1000, 2)} seconds, showing that the system was no longer operating in a stable steady state.',
        '',
        f'At 500 VUs, overload behavior became explicit. The first valid 500-VU run still delivered {fmt(valid_500[0].rps, 2)} req/s, but with P95 latency {fmt(valid_500[0].p95_ms, 2)} ms and P99 latency {fmt(valid_500[0].p99_ms, 2)} ms. The second valid run degraded to an error rate of {fmt(valid_500[1].error_rate * 100, 2)}% and a P99 of {fmt(valid_500[1].p99_ms, 2)} ms. The remaining {len(incomplete_500)} runs were incomplete, each issuing only a single failed request with zero transferred bytes. These incomplete runs indicate collapse rather than normal degraded service.',
        '',
        'The current evidence supports the following midterm conclusion: the light API path is healthy at low concurrency, reaches a clear tail-latency cliff between 10 and 50 virtual users, becomes saturated but still functional at 100 virtual users, turns unstable at 200 virtual users, and can collapse at 500 virtual users. Based on observed connection-refused failures and previous gateway restarts during the high-load runs, the API gateway is the leading bottleneck candidate, although this still needs to be confirmed using Prometheus and Grafana correlation in the final report.',
        '',
        '## Appendix A. Changes to Project Proposal',
        '',
        'The main change since the proposal is that the midterm dataset currently covers only the baseline and light API workload. The AI-intensive, mixed, spike, stress, scaling, and cache-comparison stages remain planned for the final report. In addition, high-load threshold failures required updates to the test harness so that failed runs could be preserved and the remaining batch could continue. This improved the repeatability of overload experiments and preserved evidence of failure behavior instead of aborting the campaign prematurely.',
        '',
        '## Appendix B. Updated Milestones',
        '',
        '| Date | Milestone |',
        '| --- | --- |',
        '| March 15, 2026 | Midterm report with baseline and light load-scaling results |',
        '| March 20, 2026 | Complete AI and mixed workload load-scaling experiments |',
        '| March 27, 2026 | Correlate high-load failures with Prometheus and Grafana metrics |',
        '| April 3, 2026 | Complete scaling, cache, and bottleneck-validation experiments |',
        '| April 10, 2026 | Final report submission |',
        '',
        '## Tables',
        '',
        '| Scenario | Valid runs | Mean latency (ms) | P95 latency (ms) | P99 latency (ms) | Throughput (req/s) | Error rate (%) |',
        '| --- | ---: | ---: | ---: | ---: | ---: | ---: |',
    ]

    for vus in [10, 50, 100, 200, 500]:
        agg = aggregates[vus]
        lines.append(
            f'| light / {vus} VUs | {agg.n} | {fmt(agg.avg_ms, 2)} | {fmt(agg.p95_ms, 2)} | {fmt(agg.p99_ms, 2)} | {fmt(agg.rps, 2)} | {fmt(agg.error_rate * 100, 2)} |'
        )

    lines.extend([
        '',
        '| High-load run | Incomplete | Mean latency (ms) | P95 latency (ms) | P99 latency (ms) | Throughput (req/s) | Error rate (%) |',
        '| --- | --- | ---: | ---: | ---: | ---: | ---: |',
    ])
    for vus in [200, 500]:
        for run in load_runs[vus]:
            lines.append(
                f'| {vus} / {run.file_name} | {"yes" if run.incomplete else "no"} | {fmt(run.avg_ms, 2)} | {fmt(run.p95_ms, 2)} | {fmt(run.p99_ms, 2)} | {fmt(run.rps, 2)} | {fmt(run.error_rate * 100, 2)} |'
            )

    report_path.write_text('\n'.join(lines))


def write_docx(
    docx_path: Path,
    baseline_valid: list[RunMetrics],
    load_runs: dict[int, list[RunMetrics]],
    aggregates: dict[int, AggregateMetrics],
    chart_paths: list[Path],
    host_specs: HostSpecs,
) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles['Normal'].font.name = 'Times New Roman'
    styles['Normal'].font.size = Pt(11)
    styles['Title'].font.name = 'Times New Roman'
    styles['Title'].font.size = Pt(16)
    styles['Title'].font.bold = True

    title = doc.add_paragraph()
    title.style = 'Title'
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run('P.I.T.C.H. Midterm Performance Evaluation Report')

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run('March 15, 2026').italic = True

    sections = [
        ('Abstract',
         'This midterm report presents the current performance evaluation progress for P.I.T.C.H., a microservice-based AI-assisted web application. The current dataset covers baseline characterization and light API load-scaling experiments at 10, 50, 100, 200, and 500 concurrent virtual users using k6. Initial results show that the light endpoint performs well at low load, but tail latency degrades sharply by 50 and 100 virtual users. At 200 virtual users the system becomes unstable, and at 500 virtual users the gateway enters overload behavior with partial and fully incomplete runs. These results provide a defensible midterm answer to the first project objective and identify the gateway as the leading bottleneck candidate for the final report, pending confirmation from service-level observability data.'),
        ('I. Introduction',
         'P.I.T.C.H. is a cloud-native microservice system composed of a gateway, multiple NestJS backend services, asynchronous messaging over RabbitMQ, Redis caching, persistent databases, and a Next.js frontend. The goal of this project is to evaluate how the system behaves as concurrent demand increases, determine where bottlenecks emerge, and quantify the impact of future scaling decisions.'),
        ('II. Background and Motivation',
         'Microservice architectures offer deployment flexibility and independent scaling, but they also introduce network, queuing, coordination, and observability complexity. For P.I.T.C.H., performance matters directly to user experience because requests traverse a gateway, backend services, and messaging infrastructure. The motivation for this study is to move from architecture assumptions to measured evidence and use that evidence to identify the system limits and the most valuable optimization targets.'),
        ('III. Performance Evaluation Objectives',
         '1. Measure how response time and throughput change as concurrent load increases.\n2. Identify which service or dependency becomes the bottleneck under stress.\n3. Quantify how scaling and caching affect latency, throughput, and stability in later experiments.'),
        ('IV. Evaluation Technique',
         'The current evaluation uses quantitative, measurement-based testing. Closed-workload load generation is performed with k6, and summary metrics are exported to JSON after each run. The measurements emphasized in the current dataset are mean latency, P95 latency, P99 latency, throughput, and error rate. Warm-up, steady-state, and cool-down phases are included in each run to reduce startup bias.'),
        ('V. Evaluation Setup',
         'The system under test is deployed in Docker containers. Current collected runs target the gateway entry point and exercise the light API path through the microservice stack. The broader observability environment includes Prometheus, Grafana, RabbitMQ, Redis, container monitoring, and application metrics, though this report focuses on the k6 results collected so far.'),
        ('VI. Test Workloads',
         'The current midterm dataset includes baseline characterization of the light endpoint and light API load scaling at 10, 50, 100, 200, and 500 virtual users. AI-intensive, mixed, spike, stress, replica-comparison, and cache-comparison workloads are not yet included in this report and remain part of the final-evaluation plan.'),
        ('VII. Experimental Design',
         'For the light workload, the target concurrency levels were 10, 50, 100, 200, and 500 virtual users. Five repeated runs were planned per load level. At the highest load levels, threshold failures and gateway instability emerged, so some runs represent overload behavior rather than clean steady-state service. In particular, some of the 500-virtual-user runs became incomplete and must be treated separately during analysis.'),
        ('VIII. Data Collection',
         'Data was collected from the k6 summary JSON files stored under perf/results. Two valid baseline runs were available after excluding one incomplete failed baseline. For load scaling, the light workload dataset now includes 5 runs at 10 VUs, 5 runs at 50 VUs, 4 runs at 100 VUs, 5 runs at 200 VUs, and 5 runs at 500 VUs. Of the 500-VU runs, only 2 are valid complete runs; the remaining 3 are incomplete overload runs with zero-byte transfers and should not be averaged with the valid runs.'),
    ]

    for heading, body in sections:
        doc.add_heading(heading, level=1 if heading == 'Abstract' else 2)
        for paragraph in body.split('\n'):
            doc.add_paragraph(paragraph)

    doc.add_paragraph('Host machine used for data collection:')
    for bullet in [
        f'Model: {host_specs.model_name} ({host_specs.model_identifier})',
        f'Processor: {host_specs.chip}',
        f'CPU cores: {host_specs.cpu_cores}',
        f'Memory: {host_specs.memory_gb}',
        f'Operating system: {host_specs.macos_version}',
        f'Kernel: {host_specs.kernel_version}',
    ]:
        doc.add_paragraph(bullet, style='List Bullet')

    doc.add_heading('VIII.A. Data Validity Note', level=3)
    doc.add_paragraph(
        'An important threat to validity is that the data was collected on a personal laptop that was also being used for '
        'normal student work during the testing window. Background applications, browser activity, course-related development '
        'tools, and other foreground tasks were not fully eliminated. As a result, the measured values should be interpreted '
        'as realistic development-machine results rather than fully isolated laboratory measurements. The concurrency trends '
        'and failure boundaries are still useful, but the absolute latency and throughput numbers may include noise from '
        'unrelated host activity.'
    )

    doc.add_heading('IX. Initial Results', level=2)
    doc.add_paragraph(
        f'Baseline performance was strong, with mean latency {fmt(statistics.mean(run.avg_ms for run in baseline_valid), 2)} ms, '
        f'P95 latency {fmt(statistics.mean(run.p95_ms for run in baseline_valid), 2)} ms, and throughput '
        f'{fmt(statistics.mean(run.rps for run in baseline_valid), 2)} req/s. At 10 VUs the system remained effectively identical to baseline.'
    )
    doc.add_paragraph(
        f'By 50 VUs, throughput increased to {fmt(aggregates[50].rps, 2)} req/s, but P99 latency increased to '
        f'{fmt(aggregates[50].p99_ms, 2)} ms. At 100 VUs, throughput increased further to {fmt(aggregates[100].rps, 2)} req/s '
        f'while P95 reached {fmt(aggregates[100].p95_ms, 2)} ms and P99 reached {fmt(aggregates[100].p99_ms, 2)} ms. '
        'These results indicate that tail latency degradation begins well before visible request failures.'
    )
    doc.add_paragraph(
        f'At 200 VUs, the light endpoint entered an instability regime. Across 5 runs, the aggregate error rate reached '
        f'{fmt(aggregates[200].error_rate * 100, 2)}%, mean latency reached {fmt(aggregates[200].avg_ms, 2)} ms, '
        f'and P99 latency reached {fmt(aggregates[200].p99_ms, 2)} ms. One run exhibited catastrophic behavior with '
        f'P99 latency above {fmt(load_runs[200][3].p99_ms / 1000, 2)} seconds.'
    )
    doc.add_paragraph(
        f'At 500 VUs, overload behavior became explicit. The first valid 500-VU run still delivered '
        f'{fmt([r for r in load_runs[500] if not r.incomplete][0].rps, 2)} req/s, but with P95 latency '
        f'{fmt([r for r in load_runs[500] if not r.incomplete][0].p95_ms, 2)} ms. The second valid run degraded to an '
        f'error rate of {fmt([r for r in load_runs[500] if not r.incomplete][1].error_rate * 100, 2)}% and a P99 of '
        f'{fmt([r for r in load_runs[500] if not r.incomplete][1].p99_ms, 2)} ms. The remaining 3 runs were incomplete collapse runs.'
    )
    doc.add_paragraph(
        'The current evidence supports the following midterm conclusion: the light API path is healthy at low concurrency, '
        'reaches a clear tail-latency cliff between 10 and 50 virtual users, becomes saturated but still functional at '
        '100 virtual users, turns unstable at 200 virtual users, and can collapse at 500 virtual users. Based on observed '
        'connection-refused failures and previous gateway restarts during the high-load runs, the API gateway is the leading '
        'bottleneck candidate, although this still needs to be confirmed using Prometheus and Grafana correlation in the final report.'
    )

    headers = ['Scenario', 'Valid runs', 'Mean latency (ms)', 'P95 latency (ms)', 'P99 latency (ms)', 'Throughput (req/s)', 'Error rate (%)']
    rows = []
    for vus in [10, 50, 100, 200, 500]:
        agg = aggregates[vus]
        rows.append([
            f'light / {vus} VUs',
            str(agg.n),
            fmt(agg.avg_ms, 2),
            fmt(agg.p95_ms, 2),
            fmt(agg.p99_ms, 2),
            fmt(agg.rps, 2),
            fmt(agg.error_rate * 100, 2),
        ])
    add_table(doc, headers, rows)
    add_caption(doc, 'Table I. Aggregate light-workload results across concurrent user levels.')

    for chart_path, caption in zip(
        chart_paths,
        [
            'Fig. 1. Light workload latency vs concurrency.',
            'Fig. 2. Light workload throughput and error rate vs concurrency.',
            'Fig. 3. High-load run variability for the light workload.',
        ],
    ):
        doc.add_picture(str(chart_path), width=Inches(6.6))
        add_caption(doc, caption)

    headers = ['High-load run', 'Incomplete', 'Mean latency (ms)', 'P95 latency (ms)', 'P99 latency (ms)', 'Throughput (req/s)', 'Error rate (%)']
    rows = []
    for vus in [200, 500]:
        for run in load_runs[vus]:
            rows.append([
                f'{vus} / {run.file_name}',
                'yes' if run.incomplete else 'no',
                fmt(run.avg_ms, 2),
                fmt(run.p95_ms, 2),
                fmt(run.p99_ms, 2),
                fmt(run.rps, 2),
                fmt(run.error_rate * 100, 2),
            ])
    add_table(doc, headers, rows)
    add_caption(doc, 'Table II. Per-run high-load results for 200 and 500 virtual users.')

    doc.add_heading('Appendix A. Changes to Project Proposal', level=2)
    doc.add_paragraph(
        'The main change since the proposal is that the midterm dataset currently covers only the baseline and light API workload. '
        'The AI-intensive, mixed, spike, stress, scaling, and cache-comparison stages remain planned for the final report. '
        'In addition, high-load threshold failures required updates to the test harness so that failed runs could be preserved '
        'and the remaining batch could continue.'
    )

    doc.add_heading('Appendix B. Updated Milestones', level=2)
    milestone_headers = ['Date', 'Milestone']
    milestone_rows = [
        ['March 15, 2026', 'Midterm report with baseline and light load-scaling results'],
        ['March 20, 2026', 'Complete AI and mixed workload load-scaling experiments'],
        ['March 27, 2026', 'Correlate high-load failures with Prometheus and Grafana metrics'],
        ['April 3, 2026', 'Complete scaling, cache, and bottleneck-validation experiments'],
        ['April 10, 2026', 'Final report submission'],
    ]
    add_table(doc, milestone_headers, milestone_rows)

    doc.save(docx_path)


def main() -> None:
    ensure_output_dirs()
    host_specs = get_host_specs()
    baseline_runs, baseline_valid, load_runs, aggregates = load_dataset()

    latency_chart = build_latency_chart(aggregates)
    throughput_chart = build_throughput_error_chart(aggregates)
    high_load_chart = build_high_load_run_chart(load_runs)

    markdown_path = OUTPUT_ROOT / 'pitch_midterm_report_2026-03-15.md'
    docx_path = OUTPUT_ROOT / 'pitch_midterm_report_2026-03-15.docx'

    write_markdown(markdown_path, baseline_valid, load_runs, aggregates, host_specs)
    write_docx(
        docx_path,
        baseline_valid,
        load_runs,
        aggregates,
        [latency_chart, throughput_chart, high_load_chart],
        host_specs,
    )

    summary = {
        'markdown': str(markdown_path),
        'docx': str(docx_path),
        'assets': [str(latency_chart), str(throughput_chart), str(high_load_chart)],
    }
    print(json.dumps(summary, indent=2))


if __name__ == '__main__':
    main()
